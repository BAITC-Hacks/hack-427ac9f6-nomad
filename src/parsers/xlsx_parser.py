from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.models import Document, SourceChunk
from src.services.chunking import make_chunk


def parse_xlsx(data: bytes, document: Document) -> list[SourceChunk]:
    # Keep formulas as original source text; do not calculate or guess values.
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    chunks = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
            # Do not trust incorrectly cached worksheet dimensions.
            sheet.reset_dimensions()
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                cells = []
                for column, cell in enumerate(row, start=1):
                    if cell.value is not None and str(cell.value).strip():
                        cells.append(f"{get_column_letter(column)}: {cell.value}")
                if cells:
                    chunks.append(make_chunk(
                        document,
                        f"s{sheet_index:03d}_r{row_number:06d}",
                        f"Sheet {sheet.title!r}, row {row_number}",
                        " | ".join(cells),
                    ))
    finally:
        workbook.close()
    return chunks

