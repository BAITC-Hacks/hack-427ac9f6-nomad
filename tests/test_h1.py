"""Small generated fixtures; no real organizational documents are bundled."""
from io import BytesIO
from pathlib import Path
import unittest

import pymupdf
from docx import Document as WordDocument
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from src.parsers.parser_factory import DocumentParseError, parse_document


def word_bytes(paragraphs):
    document = WordDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class ParserTests(unittest.TestCase):
    def test_pdf_page_traceability_and_stability(self):
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), "3.4 Finance")
            pdf.new_page()  # Empty pages must not renumber later pages.
            pdf.new_page().insert_text((72, 72), "9.21 Operations")
            data = pdf.tobytes()
        before = parse_document(data, "source.PDF", "BEFORE")
        repeated = parse_document(data, "source.PDF", "BEFORE")
        after = parse_document(data, "source.PDF", "AFTER")
        self.assertEqual(before, repeated)
        self.assertEqual([c.page for c in before.chunks], [1, 3])
        self.assertEqual([c.section for c in before.chunks], ["3.4", "9.21"])
        self.assertEqual(before.chunks[1].chunk_id, "before_p003_001")
        self.assertIn("9.21 Operations", before.chunks[1].text)
        self.assertEqual(before.chunks[1].document_id, before.id)
        self.assertEqual(before.chunks[1].filename, "source.PDF")
        self.assertNotEqual(before.id, after.id)
        self.assertNotEqual(before.chunks[0].chunk_id, after.chunks[0].chunk_id)

    def test_docx_keeps_original_paragraph_index(self):
        data = word_bytes(["5.3.2 Responsibility", "", "Unnumbered paragraph"])
        document = parse_document(data, "source.docx", "BEFORE")
        self.assertEqual(len(document.chunks), 2)
        self.assertEqual(document.chunks[0].section, "5.3.2")
        self.assertEqual(document.chunks[1].locator, "Body paragraph 3")
        self.assertEqual(document.chunks[1].text, "Unnumbered paragraph")
        self.assertIsNone(document.chunks[1].section)
        self.assertIsNone(document.chunks[0].page)

    def test_xlsx_keeps_sheet_row_columns_and_formulas(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Departments"
        sheet["B3"] = "Finance"
        sheet["D3"] = 0
        sheet["A5"] = "=1+1"
        workbook.create_sheet("Other")["A3"] = False
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        document = parse_document(buffer.getvalue(), "source.xlsx", "AFTER")
        self.assertEqual(len(document.chunks), 3)
        self.assertEqual(document.chunks[0].locator, "Sheet 'Departments', row 3")
        self.assertEqual(document.chunks[0].text, "B: Finance | D: 0")
        self.assertEqual(document.chunks[1].text, "A: =1+1")
        self.assertEqual(document.chunks[2].text, "A: False")
        self.assertEqual(len({c.chunk_id for c in document.chunks}), 3)

    def test_invalid_and_unsupported_uploads(self):
        for filename, data in [("file.txt", b"text"), ("file.pdf", b""),
                               ("file.pdf", b"broken"), ("file.docx", b"broken"),
                               ("file.xlsx", b"broken")]:
            with self.subTest(filename=filename, data=data):
                with self.assertRaises(DocumentParseError):
                    parse_document(data, filename, "BEFORE")

    def test_valid_documents_without_text(self):
        with pymupdf.open() as pdf:
            pdf.new_page()
            blank_pdf = pdf.tobytes()
        workbook = Workbook()
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        for filename, data in [("blank.docx", word_bytes([""])),
                               ("blank.pdf", blank_pdf),
                               ("blank.xlsx", buffer.getvalue())]:
            with self.subTest(filename=filename):
                with self.assertRaisesRegex(DocumentParseError, "No readable text"):
                    parse_document(data, filename, "BEFORE")


class AppTests(unittest.TestCase):
    def setUp(self):
        self.app = AppTest.from_file(
            str(Path(__file__).resolve().parents[1] / "app.py"),
            default_timeout=30,
        )

    def test_start_and_missing_upload_message(self):
        self.app.run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.title[0].value, "AI Org Structure Analyzer")
        self.app.button[0].click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertIn("both BEFORE and AFTER", self.app.error[0].value)

    def test_preview_from_session_documents(self):
        data = word_bytes([f"3.{i} Department {i}" for i in range(1, 8)])
        self.app.session_state["documents"] = {
            side: parse_document(data, f"{side}.docx", side)
            for side in ("BEFORE", "AFTER")
        }
        self.app.run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.code), 10)
        self.assertEqual(self.app.code[0].value, "before_para0001")
        self.assertEqual(self.app.code[5].value, "after_para0001")
        self.assertTrue(any("Total extracted chunks: 7" in x.value
                            for x in self.app.markdown))


if __name__ == "__main__":
    unittest.main()
